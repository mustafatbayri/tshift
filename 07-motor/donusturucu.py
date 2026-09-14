# -*- coding: utf-8 -*-
"""
TShift motor - PLAN DONUSTURUCU

NE YAPAR
  Musterinin Excel vardiya planlarini tek bir KANONIK tabloya cevirir.
  Ayni gun icin birden cok revizyon varsa EN SON dosyayi esas alir.

NEDEN AYRI BIR PARCA
  Dogrulayici ve cozucu, Excel'in nasil goruktugunu bilmemeli. Aralarindaki
  tek sozlesme bu kanonik tablodur. Musteri yarin baska bir sablon kullanirsa
  yalniz bu dosya degisir.

CALISTIRMA
  cd C:\\Users\\PC\\Desktop\\Tshift\\07-motor
  py donusturucu.py

GIRDI   ../06-veri/ham/plan/**/*.xlsx|xlsm
CIKTI   ../06-veri/kanonik/plan.csv      (atamalar)
        ../06-veri/kanonik/donusum.txt   (rapor: ne okundu, ne atlandi)
"""

import os
import re
import sys
import glob
import csv
import datetime
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    print("openpyxl kurulu degil:  py -m pip install openpyxl pandas")
    sys.exit(1)

KOK = os.path.dirname(os.path.abspath(__file__))
GIRDI = os.path.join(KOK, "..", "06-veri", "ham", "plan")
CIKTI = os.path.join(KOK, "..", "06-veri", "kanonik")

SAAT_RE = re.compile(r"^\s*(\d{1,2})[:.](\d{2})\s*-\s*(\d{1,2})[:.](\d{2})\s*$")
ZAMAN_RE = re.compile(r"(\d{8})_(\d{6})")

rapor = []


def yaz(s=""):
    rapor.append(str(s))
    print(s)


def dosya_zamani(yol):
    """Dosya adindaki damgadan revizyon sirasi uretir: plan_20260618_140000.xlsx"""
    m = ZAMAN_RE.search(os.path.basename(yol))
    return m.group(1) + m.group(2) if m else "00000000000000"


def basligi_bul(ws):
    """SicilNo iceren satiri bulur. Sablon degisirse burasi kirilir - kasitli."""
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
        if row and any(isinstance(v, str) and "SicilNo" in v for v in row if v):
            return r, row
    return None, None


def donustur():
    dosyalar = sorted(glob.glob(os.path.join(GIRDI, "**", "*.xls*"), recursive=True),
                      key=dosya_zamani)
    if not dosyalar:
        yaz("Plan dosyasi bulunamadi: %s" % GIRDI)
        return []

    yaz("=" * 74)
    yaz("PLAN DONUSTURUCU")
    yaz("=" * 74)
    yaz("Girdi dosyasi: %d" % len(dosyalar))
    yaz("")

    # (sicil, gun) -> kayit.  Sonra gelen dosya oncekini EZER (revizyon mantigi).
    kanonik = {}
    sayac = Counter()
    atlanan = Counter()
    kisi_bilgi = {}

    for yol in dosyalar:
        z = dosya_zamani(yol)
        try:
            wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
        except Exception as e:
            atlanan["acilamadi"] += 1
            yaz("  ACILAMADI %s -> %s" % (os.path.basename(yol), e))
            continue

        for sayfa in wb.sheetnames:
            if sayfa in ("HC", "IK"):
                sayac["atlanan_sayfa_%s" % sayfa] += 1
                continue

            ws = wb[sayfa]
            hb, bas = basligi_bul(ws)
            if hb is None:
                # CM2 sayfasi FARKLI bir sablon: SicilNo yok, kisi ADIYLA
                # anahtarlaniyor; baslangic ve bitis AYRI kolonlarda; 28 gunluk.
                # Ustelik dosyalarin cogunda tamamen #REF!.
                # Bilerek okumuyoruz - isimle anahtarlama guvenilir degil ve
                # veri zaten bozuk. Sessizce dusurmek yerine sayiyoruz.
                atlanan["CM2_farkli_sablon" if sayfa == "ÇM2" else "baslik_yok"] += 1
                continue

            idx = {v: i for i, v in enumerate(bas) if isinstance(v, str)}
            if "SicilNo" not in idx:
                atlanan["sicil_kolonu_yok"] += 1
                continue
            tarih_idx = [i for i, v in enumerate(bas) if isinstance(v, datetime.datetime)]
            if not tarih_idx:
                atlanan["tarih_kolonu_yok"] += 1
                continue

            for row in ws.iter_rows(min_row=hb + 1, values_only=True):
                sicil = row[idx["SicilNo"]] if idx["SicilNo"] < len(row) else None
                if sicil is None or not str(sicil).strip():
                    continue
                sicil = str(sicil).strip()
                if sicil.lower() in ("sicilno", "toplam"):
                    continue

                def al(ad):
                    i = idx.get(ad)
                    if i is None or i >= len(row) or row[i] is None:
                        return ""
                    return str(row[i]).strip()

                kisi_bilgi[sicil] = {
                    "ad_soyad": al("İsim & Soy İsim"),
                    "takim": al("Takım"),
                    "gorev": al("Görevi"),
                }

                for i in tarih_idx:
                    if i >= len(row):
                        continue
                    ham = row[i]
                    if ham is None:
                        continue
                    gun = bas[i].date()
                    s = str(ham).strip()
                    if not s:
                        continue

                    anahtar = (sicil, gun)
                    if anahtar in kanonik and kanonik[anahtar]["revizyon"] > z:
                        continue  # elimizde daha yeni bir revizyon var

                    m = SAAT_RE.match(s)
                    if m:
                        bsa, bda = int(m.group(1)), int(m.group(2))
                        esa, eda = int(m.group(3)), int(m.group(4))
                        b = datetime.datetime.combine(gun, datetime.time(bsa % 24, bda))
                        e = datetime.datetime.combine(gun, datetime.time(esa % 24, eda))
                        gece_asan = e <= b
                        if gece_asan:
                            e += datetime.timedelta(days=1)
                        kanonik[anahtar] = {
                            "sicil": sicil, "gun": gun, "tur": "VARDIYA",
                            "baslangic": b, "bitis": e,
                            "brut_saat": round((e - b).total_seconds() / 3600, 2),
                            "gece_asan": 1 if gece_asan else 0,
                            "ham": s, "revizyon": z,
                            "kaynak": os.path.basename(yol), "sayfa": sayfa,
                        }
                        sayac["vardiya"] += 1
                    elif re.match(r"^[\d.,\s]+$", s) or s.startswith("2026-"):
                        continue  # sayi ya da tarih artigi - atama degil
                    else:
                        kanonik[anahtar] = {
                            "sicil": sicil, "gun": gun, "tur": "DURUM",
                            "baslangic": "", "bitis": "", "brut_saat": 0,
                            "gece_asan": 0, "ham": s, "revizyon": z,
                            "kaynak": os.path.basename(yol), "sayfa": sayfa,
                        }
                        sayac["durum"] += 1
        wb.close()

    # ---- Yazma ----
    os.makedirs(CIKTI, exist_ok=True)
    satirlar = sorted(kanonik.values(), key=lambda k: (k["sicil"], k["gun"]))
    yol_csv = os.path.join(CIKTI, "plan.csv")
    with open(yol_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["sicil", "ad_soyad", "takim", "gorev", "gun", "tur",
                    "baslangic", "bitis", "brut_saat", "gece_asan",
                    "ham_deger", "revizyon", "kaynak_dosya", "sayfa"])
        for k in satirlar:
            b = kisi_bilgi.get(k["sicil"], {})
            w.writerow([
                k["sicil"], b.get("ad_soyad", ""), b.get("takim", ""), b.get("gorev", ""),
                k["gun"].isoformat(), k["tur"],
                k["baslangic"].strftime("%Y-%m-%d %H:%M") if k["baslangic"] else "",
                k["bitis"].strftime("%Y-%m-%d %H:%M") if k["bitis"] else "",
                k["brut_saat"], k["gece_asan"], k["ham"], k["revizyon"],
                k["kaynak"], k["sayfa"],
            ])

    gunler = [k["gun"] for k in satirlar]
    nihai = Counter(k["tur"] for k in satirlar)
    yaz("SONUC")
    yaz("  Kanonik kayit    : %d  (revizyonlar teklestirildi)" % len(satirlar))
    yaz("    vardiya atamasi: %d" % nihai["VARDIYA"])
    yaz("    durum kaydi    : %d" % nihai["DURUM"])
    yaz("  Okunan ham hucre : %d (ayni gun icin sonraki revizyon oncekini ezdi)"
        % (sayac["vardiya"] + sayac["durum"]))
    yaz("  Farkli calisan   : %d" % len(set(k["sicil"] for k in satirlar)))
    if gunler:
        yaz("  Tarih araligi    : %s -> %s (%d farkli gun)"
            % (min(gunler), max(gunler), len(set(gunler))))
    yaz("  Gece yarisini asan atama: %d" % sum(k["gece_asan"] for k in satirlar))
    if atlanan:
        yaz("")
        yaz("  ATLANANLAR (sablona uymayan):")
        for k, n in atlanan.most_common():
            yaz("    %-22s %d" % (k, n))
    yaz("")
    yaz("  Yazildi: %s" % os.path.normpath(yol_csv))

    with open(os.path.join(CIKTI, "donusum.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(rapor))
    return satirlar


if __name__ == "__main__":
    donustur()
