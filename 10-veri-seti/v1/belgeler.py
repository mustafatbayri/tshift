# -*- coding: utf-8 -*-
"""
TShift veri seti v1 -- BELGELER: JSON Schema, veri sozlugu, Excel import ornegi

Uc cikti da `tablolar.py`den uretilir (tek kaynak):
  sema/tshift-veri-semasi.schema.json    (JSON Schema 2020-12)
  VERI-SOZLUGU.md                        (giris metni burada; tablo bolumleri uretilir)
  excel/tshift-calisan-import-ornegi.xlsx (dogal anahtarlarla; veri/tshift-500.json'dan)

KOSTURMA
  cd C:\\Users\\PC\\Desktop\\Tshift\\10-veri-seti\\v1
  py belgeler.py
"""

import io
import json
import os
import sys

BURASI = os.path.dirname(os.path.abspath(__file__))
if BURASI not in sys.path:
    sys.path.insert(0, BURASI)

from tablolar import TABLOLAR, EXCEL_SAYFALARI, kolonlar, tablo     # noqa: E402

SEMA_YOLU = os.path.join(BURASI, "sema", "tshift-veri-semasi.schema.json")
SOZLUK_YOLU = os.path.join(BURASI, "VERI-SOZLUGU.md")
EXCEL_YOLU = os.path.join(BURASI, "excel", "tshift-calisan-import-ornegi.xlsx")

TIP = {"uuid": {"type": "string", "format": "uuid"},
       "text": {"type": "string"},
       "int": {"type": "integer"},
       "numeric": {"type": "number"},
       "bool": {"type": "boolean"},
       "date": {"type": "string", "format": "date"},
       "timestamptz": {"type": "string", "format": "date-time"},
       "int[]": {"type": "array", "items": {"type": "integer"}},
       "text[]": {"type": "array", "items": {"type": "string"}}}
KAYNAK_ADI = {"import": "import (LDAP / API / Excel)", "ayar": "firma ayarı (ekrandan)",
              "sistem": "sistem (TShift'in ürettiği)", "hesap": "hesaplanır (backend)"}


# ----------------------------------------------------------------------
# JSON Schema
# ----------------------------------------------------------------------

def _kolon_sema(tip, zorunlu, aciklama, ek):
    if tip == "json":
        s = dict(ek.get("json_sema") or {})
    else:
        s = dict(TIP[tip])
    if "enum" in ek:
        s = {"enum": list(ek["enum"])}
        if not zorunlu and None not in s["enum"]:
            s["enum"].append(None)
    elif not zorunlu:
        if "type" in s:
            s["type"] = [s["type"], "null"]
        else:
            s = {"anyOf": [s, {"type": "null"}]} if s else {}
    d = aciklama
    if ek.get("motor"):
        d += " Motor: %s." % ek["motor"]
    if ek.get("ref"):
        d += " → %s.%s" % (ek["ref"], "kod" if ek["ref"] == "rule_types" else "id")
        s["x-ref"] = ek["ref"]
    if ek.get("yeni"):
        s["x-sartnamede-yok"] = True
    s["description"] = d
    return s


def sema():
    defs = {}
    for t in TABLOLAR:
        kol = kolonlar(t)
        defs[t["ad"]] = {
            "type": "object", "title": t["ad"], "description": t["aciklama"],
            "x-bolum": t["bolum"], "x-kaynak": t["kaynak"],
            "properties": {k[0]: _kolon_sema(*k[1:]) for k in kol},
            "required": [k[0] for k in kol],             # her kolon yazilir; zorunlu olmayan null olabilir
            "additionalProperties": False}
        if t.get("yeni"):
            defs[t["ad"]]["x-sartnamede-yok"] = True
    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "$id": "https://tshift.example/sema/v1/tshift-veri-semasi.schema.json",
        "title": "TShift veri seti v1",
        "description": "Şartname §8 tabloları (M-16: kiracı alanı yok). Her tablo bir dizi; her satır tablonun "
                       "kolonlarının tamamını taşır, zorunlu olmayan kolon null olabilir. Üretim: "
                       "10-veri-seti/v1/belgeler.py (tablolar.py'den).",
        "type": "object",
        "properties": dict({"_meta": {"type": "object", "description": "Veri setinin kimliği ve sayıları."}},
                           **{t["ad"]: {"type": "array", "items": {"$ref": "#/$defs/%s" % t["ad"]}} for t in TABLOLAR}),
        "required": [t["ad"] for t in TABLOLAR],
        "additionalProperties": False,
        "$defs": defs}


# ----------------------------------------------------------------------
# Veri sozlugu (Markdown)
# ----------------------------------------------------------------------

GIRIS = """# TShift veri seti v1 — veri sözlüğü

> Yazılım ekibine (Yılmaz: LDAP / API / Excel import'u ve kuyruk) ve ürün
> tarafına. Bu dosya `belgeler.py` ile **üretilir** — elle düzeltme
> yapmayın; değişiklik `tablolar.py`'ye (tablolar) ya da `belgeler.py`'deki
> giriş metnine yazılır.

## İçindekiler

1. Bu klasör ne
2. Dosyalar
3. Veri setinin içeriği
4. Uyulan kararlar
5. Şartnameden farklar (öneri — şartnameye işlenecek)
6. Açık sorular
7. Motor girdisine dönüşüm
8. Bulgular (veri seti kurulurken çıkanlar)
9. Yeniden üretme ve test
10. Tablolar (kolon kolon)

## 1 · Bu klasör ne

Çalışan verisi ve planın bütün girdisi **şartname §8'in tablolarında**,
gerçekçi bir çağrı merkezi örneğiyle: 500 kişi, üç ekip, iki yayınlanmış
geçmiş hafta, üretilecek hafta. İsimler, e-postalar, telefonlar ve İK
kimlikleri **uydurmadır** (Mustafa, 13 Ekim: *"isimler dışında gerçekçi"*);
geri kalan her şey ölçümlerin kullandığı sahneden gelir.

İki kullanımı var:

- **Yazılım:** import'un (LDAP / API / Excel) dolduracağı tabloların biçimi
  ve örnek verisi; JSON Schema ile doğrulanır; Excel örneği aynı veriyi
  import dosyası biçiminde verir.
- **Motor testleri:** bundan sonra testler ve ölçümler bu dosyadan koşar
  (`donustur.py` motor girdisini kurar); veri her seferinde bellekte
  yeniden üretilmez (Mustafa, 13 Ekim 23:20).

## 2 · Dosyalar

| dosya | ne |
|---|---|
| `veri/tshift-500.json` | 500 kişilik veri seti (bütün tablolar, tek dosya) |
| `veri/tshift-49.json` | 49 kişilik küçük kopya (0,1 ölçek; hızlı testler) |
| `sema/tshift-veri-semasi.schema.json` | JSON Schema 2020-12 — her tablo `$defs` altında |
| `excel/tshift-calisan-import-ornegi.xlsx` | Çalışan tarafının Excel import örneği (doğal anahtarlarla) |
| `tablolar.py` | Tablo tanımları — **tek kaynak** (şema, sözlük, Excel buradan) |
| `uret.py` | Veri setini üretir (tohumlu; aynı girdiyle aynı dosya) |
| `donustur.py` | Tablolar → motor girdisi (backend'in yapacağı işin referansı) |
| `belgeler.py` | Şema, bu sözlük ve Excel'i üretir |
| `testler/test_veri_seti.py` | Bütünlük, şema, dönüştürücü ve çözülebilirlik testleri |

## 3 · Veri setinin içeriği

- **Firma:** Örnek Çağrı Merkezi A.Ş. (tek satır; M-16). Plan haftası
  **12 Ekim 2026** (Pazartesi = gün 0).
- **Organizasyon:** "Operasyon" üst departmanı; altında Satış (7/24), Back
  Office (7/24) ve Müşteri Hizmetleri (hafta içi 07–23, hafta sonu 08–19).
  Plan üst departman için üretilir ve üçünü birlikte kapsar (çift yetkinlikli
  çalışanlar iki ekibe sayılır, K-50).
- **Çalışanlar:** {calisan} kişi; {tam} tam zamanlı, {yari} yarı zamanlı;
  {pasif} pasif, {ayrildi} ayrılmış; {cift} kişi iki ekipte; {gece} kişi
  "gece çalışamaz".
- **Vardiya şablonları:** {sablon} (ekip başına sabah / ara / akşam / gece /
  uzun / kısa…), genişletilmiş saatle.
- **Kurallar:** katalogdaki 41 kuralın tamamı firmada tanımlı ({kural}
  satır; rol ve yetkinlik kapsaması satır bazlı gereklilikler).
- **Talep:** "Normal hafta (Ekim 2026)" profili, {hucre} hücre (ekip × gün ×
  saat).
- **Geçmiş:** iki yayınlanmış hafta (28 Eylül, 5 Ekim) — {atama} atama,
  {mola} mola. Bunlar motora `gecmis_vardiyalar` (gün −14…−1) olarak gider;
  Ekim'in ilk 11 gününden adalet devri (`employee_balances`) hesaplanır.
- **Kilitler:** üretilecek haftada {kilit} kilit (sabit ve yasak).

## 4 · Uyulan kararlar

| karar | veri setinde |
|---|---|
| **M-16** (11 Ekim): her müşteri ayrı konteyner | `tenant_id` hiçbir tabloda yok; `tenants` tek satır |
| **K-58** (2 Ekim) + 14 Ekim kapanışı | Sözleşme tipi yalnız `tam_zamanli` / `yari_zamanli`; kişiye saat yok; tam zamanlıda haftalık gün (5 / 6) sözleşmede. Eski "sezonluk" 13 kişi → tam zamanlı (sözleşme bitişi 31.12.2026), "stajyer" 16 kişi → yarı zamanlı |
| **K-50** çok ekipli sayım | `employee_teams` çoğul; firmada `cok_ekipli_sayim = hepsi` |
| **K-52** departman çalışma saatleri | `departments.acik` |
| **K-40** gece çalışamaz | `employees.gece_calisamaz` |
| **K-55** geçmişin kaynağı | TShift'in kendi yayınlanmış planları (`plans` / `assignments`); yıl içi fazla mesai PDKS'ten hesaplanır → `employee_balances` |
| **K-67** hedef aşımı sert | `tenant_rules`'ta `HEDEF_ASIMI` katalog türünde (SERT); merdiven tabanı referans üretmek için onu yumuşatır |
| **K-68** üst tolerans plan bazında | `plans.kural_degerleri` (örnekte boş) |

## 5 · Şartnameden farklar (öneri — şartnameye işlenecek)

Şartname §8'de **olmayan**, ama motorun okuduğu ya da kararla gelen alanlar
(şemada `x-sartnamede-yok: true`):

{farklar}

Ayrıca: `departments.acilis_saat` / `kapanis_saat` yerine `acik` (K-52; çok
pencereli saatleri iki kolon taşıyamıyor); `employee_contracts.haftalik_saat`
kaldırıldı (K-58).

## 6 · Açık sorular

1. **Plan kapsamı:** şartname planı tek departmana bağlıyor; motor üç ekibi
   birlikte çözüyor (çift yetkinlikliler yüzünden ayrı çözülemezler). Veri
   setinde plan **üst departmana** bağlı ve alt departmanları kapsıyor. Bu
   yorum doğru mu (Yılmaz)?
2. **`employee_balances`:** devir ve yıl içi fazla mesai her plan koşusunda
   mı hesaplanacak, ay sonunda mı saklanacak? Tablo önerisi bu yüzden
   "hesaplanmış ara sonuç" (Yılmaz).
3. **Kimlik tabloları** (`users` vb.) veri setinde yok; `*_user_id` kolonları
   boş. Import'un kullanıcı eşleştirmesi ayrıca konuşulmalı (Yılmaz).
4. **Kural değerlerinin plan kademesi** (T-81): şimdilik yalnız
   `ust_tolerans`; genel kural Mustafa'da.
5. **`yasak` kilidi:** şartnamede var, motor girdisi `tip: "yasak"` ile
   alıyor; doğrulayıcı tarafı ayrıca bakılmalı.

## 7 · Motor girdisine dönüşüm

`donustur.py` — backend'in plan koşusu için yapacağı işin referans
uygulaması; kuyruğa konacak iş mesajı bu çıktıdır (şartname §11.2):

- **Kapsam:** planın departmanı + bütün alt departmanları. Çalışan,
  departmanı kapsamdaysa gider; durumu ne olursa olsun (pasif / ayrılmış
  kişiyi `AKTIF_CALISAN` dışlar).
- **Sözleşme:** plan haftasının ilk günü geçerli olan; yoksa en son
  başlayan. Tam zamanlı → `haftalik_saat: 45` (kanun), `gun_sayisi =
  haftalik_gun`.
- **İzin:** plan haftasına düşen **tüm gün** izinler, gün gün (K-31).
- **Uygunluk:** haftalık satırlar aynen; tarihli satırlar haftaya açılır.
- **Geçmiş:** plandan önceki iki yayınlanmış haftanın seçili versiyonu →
  `gecmis_vardiyalar` (molalarıyla) ve `gecmis_bilinen_gunler` (−14…−1).
- **Kurallar:** aktif `tenant_rules`; tür = geçersiz kılma ya da katalog;
  yasal / kabul edilebilir **katalogdan**; satır bazlı kuralın her
  gerekliliği ayrı kural; planın `kural_degerleri` parametrelere eklenir.
- **Ağırlıklar:** varsayılan profilin ağırlıkları.

Bugünkü ölçüm sahnesiyle (`merdiven.taban_sahne`) farkı yalnız bilerek
değişen şeylerdir (test `test_DONUSTURUCU_eski_sahneden_YALNIZ_bilinen_farklar`):
29 kişinin sözleşmesi, geçmiş (yeni), devir (geçmişten), 6 kuralın
bayrakları (katalogdan, §8), `HEDEF_ASIMI` türü, `sektor`, departman adları,
`agirliklar` (yeni anahtar).

## 8 · Bulgular (veri seti kurulurken çıkanlar)

- **Üreticinin kural bayrakları katalogdan sapmış** (6 kural): {sapmalar}.
  Veri setinde katalog geçerli (sistem tablosu, K-18 / K-20). Motorun
  planını değiştirmez; yayın kapısının "kabul edilebilir" ve "yasal"
  cümlelerini değiştirir.
- **Geçmiş + sert hedef aşımı:** geçen haftanın Pazar gecesi vardiyası
  Pazartesi sabahını kapatıyor (K-56); bu talep eğrisiyle (hedef 5,
  sahada 7) `HEDEF_ASIMI` sertken plan **ön kontrolde çözümsüz** — talep
  geçmişi bilmeden girilmiş. Merdiven tabanı (aşım yumuşak) sorunsuz
  çözülüyor; seviye 1 talebi yeni referanstan türeyince tutarlı olur.

## 9 · Yeniden üretme ve test

```powershell
cd C:\\Users\\PC\\Desktop\\Tshift\\10-veri-seti\\v1
py uret.py        # veri/tshift-500.json, veri/tshift-49.json
py belgeler.py    # şema, bu sözlük, Excel
py -m pytest testler -q
```

## 10 · Tablolar (kolon kolon)

Her tabloda (kural kataloğu hariç) ortak kolonlar: `id` (UUID, birincil
anahtar), `created_at`, `updated_at`. Aşağıda **Z** = zorunlu (boş olamaz),
**→** = ilişki (yabancı anahtar), 🆕 = şartnamede yok.
"""


def _sayilar(veri):
    T = veri
    soz = {s["employee_id"]: s for s in T["employee_contracts"]}
    uye = {}
    for u in T["employee_teams"]:
        uye[u["employee_id"]] = uye.get(u["employee_id"], 0) + 1
    return {"calisan": len(T["employees"]),
            "tam": sum(1 for s in soz.values() if s["tip"] == "tam_zamanli"),
            "yari": sum(1 for s in soz.values() if s["tip"] == "yari_zamanli"),
            "pasif": sum(1 for e in T["employees"] if e["durum"] == "pasif"),
            "ayrildi": sum(1 for e in T["employees"] if e["durum"] == "ayrildi"),
            "cift": sum(1 for n in uye.values() if n > 1),
            "gece": sum(1 for e in T["employees"] if e["gece_calisamaz"]),
            "sablon": len(T["shift_templates"]), "kural": len(T["tenant_rules"]),
            "hucre": len(T["demand_cells"]), "atama": len(T["assignments"]),
            "mola": len(T["assignment_breaks"]), "kilit": len(T["plan_locks"])}


def _md(metin):
    return str(metin).replace("|", "\\|").replace("\n", " ")


def sozluk(veri):
    farklar = []
    for t in TABLOLAR:
        if t.get("yeni"):
            farklar.append("- **`%s`** (tablo) — %s" % (t["ad"], t["aciklama"].split(". ")[0] + "."))
        for k in t["kolonlar"]:
            if k[4].get("yeni"):
                farklar.append("- `%s.%s` — %s" % (t["ad"], k[0], k[3]))
    sapma = veri["_meta"]["donusum"]["katalogdan_sapan_uretici_bayraklari"]
    metin = GIRIS.format(farklar="\n".join(farklar),
                         sapmalar="; ".join("`%s`" % s.split(":")[0] for s in sapma),
                         **_sayilar(veri))
    # tablolar
    bolum = None
    sira = 0
    for t in TABLOLAR:
        if t["bolum"] != bolum:
            bolum = t["bolum"]
            metin += "\n### %s\n" % bolum
        sira += 1
        metin += "\n#### `%s`%s\n\n" % (t["ad"], " 🆕" if t.get("yeni") else "")
        metin += "%s\n\n" % t["aciklama"]
        metin += "Kaynak: **%s** · Örnekte %d satır.\n\n" % (KAYNAK_ADI[t["kaynak"]], len(veri[t["ad"]]))
        metin += "| kolon | tip | Z | açıklama | motor girdisi |\n|---|---|---|---|---|\n"
        for ad, tip, zorunlu, aciklama, ek in (t["kolonlar"] if t.get("ortak_yok") else t["kolonlar"]):
            tip_y = tip + ((" → `%s`" % ek["ref"]) if ek.get("ref") else "")
            ack = aciklama
            if "enum" in ek:
                ack += " Değerler: " + ", ".join("`%s`" % ("boş" if e is None else e) for e in ek["enum"]) + "."
            metin += "| `%s`%s | %s | %s | %s | %s |\n" % (
                ad, " 🆕" if ek.get("yeni") else "", tip_y, "Z" if zorunlu else "", _md(ack),
                ("`%s`" % _md(ek["motor"])) if ek.get("motor") else "")
    return metin


# ----------------------------------------------------------------------
# Excel import ornegi
# ----------------------------------------------------------------------

def excel(veri, yol=EXCEL_YOLU):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    kod = {ad: {s["id"]: s.get("kod") for s in veri[ad]} for ad in ("departments", "teams", "operational_roles", "skills")}
    pno = {e["id"]: e["personel_no"] for e in veri["employees"]}
    font, baslik_font = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    baslik_dolgu = PatternFill("solid", fgColor="3D5A80")
    zorunlu_dolgu = PatternFill("solid", fgColor="1F3A5F")

    wb = Workbook()
    ack = wb.active
    ack.title = "Aciklama"
    satirlar = [
        ("TShift — çalışan import örneği (veri seti v1)", True),
        ("Bu dosya 10-veri-seti/v1/veri/tshift-500.json'un çalışan tarafıdır; belgeler.py ile üretilir.", False),
        ("İsimler, e-posta, telefon ve İK kimlikleri UYDURMADIR.", False),
        ("", False),
        ("Nasıl doldurulur", True),
        ("• Her sayfa bir tablodur; ilk satır kolon adıdır. Koyu mavi başlık = zorunlu, açık mavi = isteğe bağlı.", False),
        ("• Çalışan her sayfada personel_no ile bağlanır; ekip, departman, rol ve yetkinlik KOD ile yazılır "
         "(kodlar 'Kodlar' sayfasında). Sistem kimlikleri (UUID) import'ta üretilir.", False),
        ("• Tarih: YYYY-AA-GG. Saat: ondalıklı (8,5 = 08:30); gece yarısını aşan bitiş 24'ten büyük olabilir.", False),
        ("• Doğru / yanlış: DOĞRU / YANLIŞ (Excel mantıksal değeri).", False),
        ("• Başlık hücresinin notunda kolonun açıklaması ve izin verilen değerler var.", False),
        ("• Örnek satırlar dolu: biçimi göstermek için; kendi verinizle değiştirin.", False),
        ("", False),
        ("Sayfalar", True),
    ]
    for sayfa, tablo_adi, _ in EXCEL_SAYFALARI:
        satirlar.append(("• %s — %s" % (sayfa, tablo(tablo_adi)["aciklama"].replace("**", "")), False))
    satirlar.append(("• Kodlar — departman, ekip, rol ve yetkinlik kodları.", False))
    for i, (m, kalin) in enumerate(satirlar, 1):
        c = ack.cell(row=i, column=1, value=m)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=kalin)
    ack.column_dimensions["A"].width = 140

    for sayfa, tablo_adi, eslem in EXCEL_SAYFALARI:
        ws = wb.create_sheet(sayfa)
        t = tablo(tablo_adi)
        kmeta = {k[0]: k for k in kolonlar(t)}
        for j, (excel_kol, kaynak) in enumerate(eslem, 1):
            if kaynak == "calisan":
                meta = ("personel_no", "text", True, "Çalışanın sicil no'su (Calisanlar sayfasındaki).", {})
            elif kaynak.startswith("ref:"):
                hedef = kaynak.split(":")[2]
                meta = kmeta[hedef]
                meta = (excel_kol, "text", meta[2], meta[3] + " (KOD ile)", {k: v for k, v in meta[4].items() if k != "ref"})
            else:
                meta = kmeta[kaynak]
            c = ws.cell(row=1, column=j, value=excel_kol)
            c.font = baslik_font
            c.fill = zorunlu_dolgu if meta[2] else baslik_dolgu
            c.alignment = Alignment(vertical="center")
            not_ = meta[3].replace("**", "").replace("`", "")
            if "enum" in meta[4]:
                not_ += "\nDeğerler: " + ", ".join("boş" if e is None else str(e) for e in meta[4]["enum"])
                degerler = [str(e) for e in meta[4]["enum"] if e is not None]
                dv = DataValidation(type="list", formula1='"%s"' % ",".join(degerler), allow_blank=not meta[2])
                ws.add_data_validation(dv)
                dv.add("%s2:%s5000" % (get_column_letter(j), get_column_letter(j)))
            c.comment = Comment(("ZORUNLU. " if meta[2] else "") + not_, "TShift")
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(excel_kol) + 4))
        for i, satir in enumerate(veri[tablo_adi], 2):
            for j, (excel_kol, kaynak) in enumerate(eslem, 1):
                if kaynak == "calisan":
                    v = pno[satir["employee_id"]]
                elif kaynak.startswith("ref:"):
                    ref_tablo, _, alan = kaynak[4:].partition(".")
                    hedef = alan.split(":")[1]
                    v = kod[ref_tablo].get(satir[hedef]) if satir[hedef] else None
                else:
                    v = satir[kaynak]
                    if isinstance(v, list):
                        v = ",".join(str(x) for x in v)
                c = ws.cell(row=i, column=j, value=v)
                c.font = font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(eslem)), max(2, len(veri[tablo_adi]) + 1))

    ws = wb.create_sheet("Kodlar")
    r = 1
    for baslik, ad, alanlar in (("Departmanlar", "departments", ("kod", "ad")), ("Ekipler", "teams", ("kod", "ad")),
                                ("Roller", "operational_roles", ("kod", "ad")), ("Yetkinlikler", "skills", ("kod", "ad"))):
        c = ws.cell(row=r, column=1, value=baslik)
        c.font = Font(name="Arial", size=10, bold=True)
        r += 1
        for s in veri[ad]:
            for j, alan in enumerate(alanlar, 1):
                ws.cell(row=r, column=j, value=s[alan]).font = font
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 32
    wb.save(yol)
    return yol


def yaz():
    v = json.load(io.open(os.path.join(BURASI, "veri", "tshift-500.json"), encoding="utf-8"))
    with io.open(SEMA_YOLU, "w", encoding="utf-8", newline="\n") as f:
        json.dump(sema(), f, ensure_ascii=False, indent=1)
        f.write("\n")
    with io.open(SOZLUK_YOLU, "w", encoding="utf-8", newline="\n") as f:
        f.write(sozluk(v))
    excel(v)
    return SEMA_YOLU, SOZLUK_YOLU, EXCEL_YOLU


if __name__ == "__main__":
    for y in yaz():
        print(os.path.relpath(y, BURASI))
